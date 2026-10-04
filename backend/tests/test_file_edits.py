from __future__ import annotations

import hashlib
import io
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import select, text

from pixel_station.app import create_app
from pixel_station.database import Attachment, DocumentChunk
from pixel_station.file_edits import (
    LEGACY_TARGETED_SCOPE,
    FileEditProposal,
    FileRevision,
    PlannedTargetedEdit,
    create_file_edit_router,
    propose_chat_edit,
)


class EditLLM:
    def __init__(self):
        self.file_id = ""

    async def models(self):
        return {"available": True, "models": []}

    async def structured(self, model, messages, schema, **kwargs):
        return schema(file_id=self.file_id, plan="Update the deadline and retain the budget.",
                      content="Deadline: Friday\nBudget: 100")


@pytest.fixture
def app(tmp_path):
    app = create_app(tmp_path, llm=EditLLM(), discover=False)
    if not hasattr(app.state, "file_edits"):
        app.include_router(create_file_edit_router(app))
    return app


@pytest.fixture
def client(app):
    with TestClient(app) as client:
        yield client


def upload(client, filename="notes.txt", content=b"Deadline: Monday\nBudget: 100"):
    response = client.post("/api/files/upload", files={"file": (filename, content)})
    assert response.status_code == 200, response.text
    return response.json()


def propose(client, file_id, content="Deadline: Friday\nBudget: 100"):
    response = client.post(f"/api/files/{file_id}/edit-proposals",
                           json={"content": content, "plan": "Move the deadline to Friday."})
    assert response.status_code == 200, response.text
    return response.json()


def confirm(client, proposal_id):
    return client.post(f"/api/files/edit-proposals/{proposal_id}/confirm", json={"confirmed": True})


def test_edit_requires_confirmation_preserves_original_and_rebuilds_search(client, app):
    original = upload(client)
    original_bytes = Path(original["path"]).read_bytes()
    before_chunks = client.get(f"/api/files/{original['id']}").json()["chunks"]
    proposal = propose(client, original["id"])
    assert client.get(f"/api/files/{original['id']}/content").content == original_bytes
    assert proposal["before_sha256"] == original["sha256"]
    assert proposal["preview_content"] == "Deadline: Friday\nBudget: 100"
    assert confirm(client, proposal["id"]).status_code == 200
    assert confirm(client, proposal["id"]).status_code == 409
    edited = client.get(f"/api/files/{original['id']}").json()
    assert edited["id"] == original["id"]
    assert edited["sha256"] != original["sha256"]
    assert edited["parser"] == "native-edit"
    assert edited["parse_status"] == "ready"
    assert edited["chunks"][0]["text"] == "Deadline: Friday\nBudget: 100"
    assert edited["chunks"][0]["id"] != before_chunks[0]["id"]
    assert edited["chunks"][0]["embedding"] is None
    assert Path(original["path"]).read_bytes() == original_bytes
    revisions = client.get(f"/api/files/{original['id']}/revisions").json()
    assert len(revisions) == 1
    assert "before_path" not in revisions[0]
    assert client.get(f"/api/files/{original['id']}/revisions/{revisions[0]['id']}/content").content == original_bytes
    with app.state.database.session() as session:
        assert session.execute(text("SELECT COUNT(*) FROM document_chunks_fts WHERE document_chunks_fts MATCH 'Friday'")).scalar() == 1
        assert session.execute(text("SELECT COUNT(*) FROM document_chunks_fts WHERE document_chunks_fts MATCH 'Monday'")).scalar() == 0


@pytest.mark.parametrize("format,content", [
    ("txt", "Revised plain text"), ("md", "# Revised heading\n\nParagraph"),
    ("csv", "name,amount\nAlice,23"), ("xlsx", "name,amount\nAlice,23"),
    ("docx", "Revised title\nA second paragraph"), ("pdf", "Revised title\nA second paragraph")
])
def test_edit_supported_export_formats_create_valid_replacements(client, format, content):
    original = client.post("/api/files/create", json={"filename": "report", "format": format,
                                                    "content": "Original title\nSecond paragraph"}).json()
    editable = client.get(f"/api/files/{original['id']}/edit-content")
    assert editable.status_code == 200, editable.text
    proposal = propose(client, original["id"], content)
    response = confirm(client, proposal["id"])
    assert response.status_code == 200, response.text
    reread = client.get(f"/api/files/{original['id']}/edit-content").json()
    assert "Revised" in reread["content"] or "Alice" in reread["content"]
    assert client.get(f"/api/files/{original['id']}/content").content


def test_xlsx_edit_preserves_other_sheets_styles_and_blocks_formula_input(client):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    workbook = Workbook()
    workbook.active.title = "Costs"
    workbook.active.append(["Name", "Cost"])
    workbook.active.append(["Alice", 10])
    workbook.active["A1"].font = Font(bold=True)
    workbook.create_sheet("Summary")["A1"] = "=SUM(Costs!B2:B10)"
    output = io.BytesIO()
    workbook.save(output)
    original = upload(client, "costs.xlsx", output.getvalue())
    proposal = propose(client, original["id"], "Name,Cost\n=HYPERLINK(1),20")
    assert "first worksheet 'Costs'" in proposal["scope"]
    assert confirm(client, proposal["id"]).status_code == 200
    updated = load_workbook(io.BytesIO(client.get(f"/api/files/{original['id']}/content").content))
    assert updated.sheetnames == ["Costs", "Summary"]
    assert updated["Summary"]["A1"].value == "=SUM(Costs!B2:B10)"
    assert updated["Costs"]["A1"].font.bold
    assert updated["Costs"]["A2"].value == "'=HYPERLINK(1)"
    assert updated["Costs"]["B2"].value == 20
    assert updated["Costs"]["B2"].data_type == "n"


def test_xlsx_review_preserves_unchanged_types_and_updates_typed_cells(client):
    from datetime import datetime

    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Amount", "Active", "Date", "Formula", "Text", "New"])
    sheet.append([10, True, datetime(2099, 10, 16), "=A2*2", "001", None])
    sheet["A2"].number_format = "0.00"
    sheet["A2"].font = Font(italic=True)
    output = io.BytesIO()
    workbook.save(output)
    original = upload(client, "typed.xlsx", output.getvalue())
    current = client.get(f"/api/files/{original['id']}/edit-content").json()
    assert "Unchanged cells" in current["warning"]
    content = current["content"].replace("10,True,2099-10-16 00:00:00", "-25.5,False,2099-10-18")
    content = content.replace("001,", "001,=DANGEROUS(1)")
    proposal = propose(client, original["id"], content)
    assert client.get(f"/api/files/{original['id']}/content").content == output.getvalue()
    assert confirm(client, proposal["id"]).status_code == 200
    updated = load_workbook(io.BytesIO(client.get(f"/api/files/{original['id']}/content").content))
    row = updated.active
    assert row["A2"].value == -25.5 and row["A2"].data_type == "n"
    assert row["A2"].number_format == "0.00" and row["A2"].font.italic
    assert row["B2"].value is False and row["B2"].data_type == "b"
    assert row["C2"].value == datetime(2099, 10, 18)
    assert row["D2"].value == "=A2*2" and row["D2"].data_type == "f"
    assert row["E2"].value == "001" and row["E2"].data_type == "s"
    assert row["F2"].value == "'=DANGEROUS(1)" and row["F2"].data_type == "s"


@pytest.mark.parametrize("replacement", ["not a number", "=SUM(1)", "1e999"])
def test_xlsx_rejects_invalid_numeric_type_without_changing_source(client, replacement):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active["A1"] = 23
    output = io.BytesIO()
    workbook.save(output)
    original = upload(client, "number.xlsx", output.getvalue())
    response = client.post(f"/api/files/{original['id']}/edit-proposals", json={"content": replacement})
    assert response.status_code == 422
    assert "finite number" in response.json()["detail"]
    assert client.get(f"/api/files/{original['id']}/content").content == output.getvalue()


def test_docx_indexes_header_footer_variants_and_tables_once(client):
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.shared import Inches

    document = Document()
    document.add_paragraph("Body reference")
    section = document.sections[0]
    section.header.paragraphs[0].text = "Shared header reference"
    section.first_page_header.paragraphs[0].text = "First page header reference"
    section.even_page_header.paragraphs[0].text = "Even page header reference"
    section.footer.paragraphs[0].text = "Shared footer reference"
    section.first_page_footer.add_table(rows=1, cols=1, width=Inches(3)).cell(0, 0).text = "Footer table reference"
    document.add_section(WD_SECTION_START.NEW_PAGE)
    output = io.BytesIO()
    document.save(output)
    original = upload(client, "stories.docx", output.getvalue())
    chunks = client.get(f"/api/files/{original['id']}").json()["chunks"]
    texts = "\n".join(chunk["text"] for chunk in chunks)
    assert texts.count("Shared header reference") == 1
    assert texts.count("Shared footer reference") == 1
    assert "First page header reference" in texts and "Even page header reference" in texts
    assert "Footer table reference" in texts
    assert any("header" in chunk["location"] for chunk in chunks)
    assert any("footer" in chunk["location"] for chunk in chunks)


def checkbox_docx(*, locked=False, inconsistent=False):
    from docx import Document
    from lxml import etree

    from pixel_station.docx_controls import C, W

    document = Document(io.BytesIO(contract_docx()))
    paragraph = document.add_paragraph("Sound service ")
    control = etree.SubElement(paragraph._p, W + "sdt")
    properties = etree.SubElement(control, W + "sdtPr")
    etree.SubElement(properties, W + "alias").set(W + "val", "Sound system")
    if locked:
        etree.SubElement(properties, W + "lock").set(W + "val", "contentLocked")
    checkbox = etree.SubElement(properties, C + "checkbox")
    etree.SubElement(checkbox, C + "checked").set(C + "val", "0")
    for name, value in (("checkedState", "2612"), ("uncheckedState", "2610")):
        state = etree.SubElement(checkbox, C + name)
        state.set(C + "val", value)
        state.set(C + "font", "MS Gothic")
    content = etree.SubElement(control, W + "sdtContent")
    run = etree.SubElement(content, W + "r")
    etree.SubElement(run, W + "t").text = "☒" if inconsistent else "☐"
    paragraph = document.add_paragraph("Cleaning service ")
    run = etree.SubElement(paragraph._p, W + "r")
    begin = etree.SubElement(run, W + "fldChar")
    begin.set(W + "fldCharType", "begin")
    data = etree.SubElement(begin, W + "ffData")
    etree.SubElement(data, W + "name").set(W + "val", "Cleaning")
    checkbox = etree.SubElement(data, W + "checkBox")
    etree.SubElement(checkbox, W + "default").set(W + "val", "1")
    run = etree.SubElement(paragraph._p, W + "r")
    etree.SubElement(run, W + "instrText").text = " FORMCHECKBOX "
    run = etree.SubElement(paragraph._p, W + "r")
    etree.SubElement(run, W + "fldChar").set(W + "fldCharType", "separate")
    run = etree.SubElement(paragraph._p, W + "r")
    etree.SubElement(run, W + "t").text = "☒"
    run = etree.SubElement(paragraph._p, W + "r")
    etree.SubElement(run, W + "fldChar").set(W + "fldCharType", "end")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_reviewed_word_checkbox_states_preserve_unrelated_parts_and_reindex(client):
    source = checkbox_docx()
    original = upload(client, "form.docx", source)
    targets = client.get(f"/api/files/{original['id']}/edit-targets").json()
    controls = targets["checkboxes"]
    assert {control["kind"] for control in controls} == {"content_control", "legacy_field"}
    assert {control["label"]: control["checked"] for control in controls} == {"Sound system": False, "Cleaning": True}
    response = client.post(f"/api/files/{original['id']}/targeted-edit-proposals", json={
        "before_sha256": original["sha256"], "checkbox_changes": [
            {"location": control["location"], "before": control["checked"], "after": not control["checked"]}
            for control in controls]})
    assert response.status_code == 200, response.text
    proposal = response.json()
    assert proposal["changes"] == [] and len(proposal["checkbox_changes"]) == 2
    assert "Before: unchecked" in proposal["preview_content"] and "After: checked" in proposal["preview_content"]
    preview = client.get(f"/api/files/edit-proposals/{proposal['id']}/preview").content
    with zipfile.ZipFile(io.BytesIO(source)) as before, zipfile.ZipFile(io.BytesIO(preview)) as after:
        assert before.namelist() == after.namelist()
        assert [name for name in before.namelist() if before.read(name) != after.read(name)] == ["word/document.xml"]
    assert client.get(f"/api/files/{original['id']}/content").content == source
    assert confirm(client, proposal["id"]).status_code == 200
    assert confirm(client, proposal["id"]).status_code == 409
    updated = client.get(f"/api/files/{original['id']}/edit-targets").json()
    assert {control["label"]: control["checked"] for control in updated["checkboxes"]} == {"Sound system": True, "Cleaning": False}
    chunks = client.get(f"/api/files/{original['id']}").json()["chunks"]
    text = "\n".join(chunk["text"] for chunk in chunks)
    assert "Sound system: checked" in text and "Cleaning: unchecked" in text


@pytest.mark.parametrize("modification, expected", [
    ("stale", 409), ("duplicate", 422), ("unknown", 422), ("same", 422), ("string-state", 422),
])
def test_checkbox_proposals_reject_unreviewed_or_ambiguous_states(client, modification, expected):
    source = checkbox_docx()
    original = upload(client, "form.docx", source)
    control = client.get(f"/api/files/{original['id']}/edit-targets").json()["checkboxes"][0]
    change = {"location": control["location"], "before": False, "after": True}
    if modification == "stale":
        change["before"] = True
    elif modification == "unknown":
        change["location"] = "word/document.xml:drawn-rectangle:0"
    elif modification == "same":
        change["after"] = False
    elif modification == "string-state":
        change["after"] = "true"
    response = client.post(f"/api/files/{original['id']}/targeted-edit-proposals", json={
        "before_sha256": original["sha256"],
        "checkbox_changes": [change, change] if modification == "duplicate" else [change]})
    assert response.status_code == expected
    assert client.get(f"/api/files/{original['id']}/content").content == source


@pytest.mark.parametrize("option", ["locked", "inconsistent"])
def test_locked_or_inconsistent_controls_are_not_exposed_for_toggling(client, option):
    original = upload(client, "limited-form.docx", checkbox_docx(**{option: True}))
    targets = client.get(f"/api/files/{original['id']}/edit-targets").json()
    assert [control["kind"] for control in targets["checkboxes"]] == ["legacy_field"]
    assert targets["unsupported_checkbox_count"] == 1


def test_text_replacements_cannot_bypass_checkbox_state_review(client):
    original = upload(client, "form.docx", checkbox_docx())
    response = targeted(client, original, [("Sound service", "☐", "☒")])
    assert response.status_code == 422
    assert "reviewed checkbox change" in response.json()["detail"]


def test_docx_edit_preserves_tables_and_paragraph_styles(client):
    from docx import Document
    document = Document()
    document.add_paragraph("Original heading", style="Heading 1")
    document.add_paragraph("Existing body")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Keep this table"
    output = io.BytesIO()
    document.save(output)
    original = upload(client, "document.docx", output.getvalue())
    proposal = propose(client, original["id"], "Updated heading\nUpdated body")
    assert "Inline formatting" in proposal["plan"]
    assert confirm(client, proposal["id"]).status_code == 200
    updated = Document(io.BytesIO(client.get(f"/api/files/{original['id']}/content").content))
    assert updated.paragraphs[0].text == "Updated heading"
    assert updated.paragraphs[0].style.name == "Heading 1"
    assert updated.tables[0].cell(0, 0).text == "Keep this table"


def test_invalid_csv_edit_leaves_original_and_index_unchanged(client, app):
    original = upload(client, "costs.csv", b"name,amount\nAlice,10")
    response = client.post(f"/api/files/{original['id']}/edit-proposals", json={"content": 'name,amount\n"unclosed'})
    assert response.status_code == 422
    assert client.get(f"/api/files/{original['id']}/content").content == b"name,amount\nAlice,10"
    with app.state.database.session() as session:
        assert list(session.scalars(select(FileEditProposal))) == []
        assert list(session.scalars(select(FileRevision))) == []


def test_stale_proposal_after_another_revision_cannot_apply(client):
    original = upload(client)
    first, second = propose(client, original["id"]), propose(client, original["id"], "Competing edit")
    assert confirm(client, first["id"]).status_code == 200
    stale = confirm(client, second["id"])
    assert stale.status_code == 409
    assert "changed after review" in stale.text
    assert client.get(f"/api/files/{original['id']}/content").content == b"Deadline: Friday\nBudget: 100"


def test_expired_rejected_and_missing_confirmation_do_not_write(client, app):
    original = upload(client)
    proposal = propose(client, original["id"])
    endpoint = f"/api/files/edit-proposals/{proposal['id']}/confirm"
    assert client.post(endpoint, json={"confirmed": False}).status_code == 422
    assert client.post(endpoint, json={"confirmed": True, "content": "unreviewed"}).status_code == 422
    assert client.delete(f"/api/files/edit-proposals/{proposal['id']}").status_code == 200
    assert confirm(client, proposal["id"]).status_code == 409
    app.state.file_edits.ttl = -1
    expired = propose(client, original["id"])
    assert confirm(client, expired["id"]).status_code == 409
    assert client.get(f"/api/files/{original['id']}/content").content == b"Deadline: Monday\nBudget: 100"


@pytest.mark.parametrize("alteration", ["proposal", "stage", "original"])
def test_mutating_reviewed_state_blocks_confirmation(client, app, alteration):
    original = upload(client)
    proposal = propose(client, original["id"])
    if alteration == "proposal":
        with app.state.database.session() as session:
            row = session.get(FileEditProposal, proposal["id"])
            row.content = "Unreviewed content"
            session.commit()
    elif alteration == "stage":
        app.state.file_edits._stage(proposal["id"], ".txt").write_text("Unreviewed content")
    else:
        Path(original["path"]).write_bytes(b"External modification")
    assert confirm(client, proposal["id"]).status_code == 409
    with app.state.database.session() as session:
        assert session.get(Attachment, original["id"]).sha256 == original["sha256"]
        assert list(session.scalars(select(FileRevision))) == []


def test_concurrent_confirmation_is_single_use(client, app):
    original = upload(client)
    proposal = propose(client, original["id"])

    def apply():
        try:
            app.state.file_edits.confirm(proposal["id"])
            return 200
        except HTTPException as exc:
            return exc.status_code

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: apply(), range(2)))
    assert sorted(results) == [200, 409]
    assert len(client.get(f"/api/files/{original['id']}/revisions").json()) == 1


def test_edit_does_not_follow_an_outside_managed_file_path(client, app, tmp_path):
    original = upload(client)
    outside = tmp_path.parent / ("outside-" + original["id"] + ".txt")
    outside.write_text("Outside data")
    try:
        with app.state.database.session() as session:
            row = session.get(Attachment, original["id"])
            row.path = str(outside)
            row.sha256 = hashlib.sha256(outside.read_bytes()).hexdigest()
            session.commit()
        assert client.get(f"/api/files/{original['id']}/edit-content").status_code == 404
        assert client.post(f"/api/files/{original['id']}/edit-proposals", json={"content": "Change outside"}).status_code == 404
        assert outside.read_text() == "Outside data"
    finally:
        outside.unlink()


def test_failed_revalidation_rolls_back_library_change(client, app, monkeypatch):
    original = upload(client)
    proposal = propose(client, original["id"])

    def fail(*args, **kwargs):
        raise ValueError("Revalidation failed")

    monkeypatch.setattr("pixel_station.file_edits.NativeFileParser.parse", fail)
    assert confirm(client, proposal["id"]).status_code == 422
    with app.state.database.session() as session:
        assert session.get(Attachment, original["id"]).sha256 == original["sha256"]
        assert session.get(FileEditProposal, proposal["id"]).status == "pending"
        assert list(session.scalars(select(FileRevision))) == []
        assert "Monday" in session.scalar(select(DocumentChunk)).text


def test_proposal_and_revision_persist_after_app_restart(tmp_path):
    app = create_app(tmp_path, llm=EditLLM(), discover=False)
    if not hasattr(app.state, "file_edits"):
        app.include_router(create_file_edit_router(app))
    with TestClient(app) as client:
        original = upload(client)
        proposal = propose(client, original["id"])
    restarted = create_app(tmp_path, llm=EditLLM(), discover=False)
    if not hasattr(restarted.state, "file_edits"):
        restarted.include_router(create_file_edit_router(restarted))
    with TestClient(restarted) as client:
        assert confirm(client, proposal["id"]).status_code == 200
        assert len(client.get(f"/api/files/{original['id']}/revisions").json()) == 1
        assert client.get(f"/api/files/{original['id']}/content").content == b"Deadline: Friday\nBudget: 100"


@pytest.mark.asyncio
async def test_chat_edit_returns_proposal_with_exact_target_without_writing(app):
    from pixel_station.files import ingest
    with app.state.database.session() as session:
        original = ingest(session, app.state.data_dir, "notes.txt", b"Deadline: Monday\nBudget: 100")
    app.state.llm.file_id = original.id
    settings = app.state.settings()
    settings.roles["planner"] = "test-local"
    app.state.set_settings(settings)
    result = await propose_chat_edit(app, [original.id], "Move the deadline to Friday")
    assert result["file_edit"]["file_id"] == original.id
    assert "confirm" in result["content"]
    assert Path(original.path).read_bytes() == b"Deadline: Monday\nBudget: 100"
    with pytest.raises(HTTPException) as error:
        await propose_chat_edit(app, [], "Edit something")
    assert error.value.status_code == 422


def test_remote_origin_cannot_propose_file_edits(client):
    original = upload(client)
    response = client.post(f"/api/files/{original['id']}/edit-proposals", json={"content": "Updated"},
                           headers={"Origin": "https://example.com"})
    assert response.status_code == 403


def contract_docx():
    from docx import Document
    from docx.shared import Inches
    from PIL import Image

    document = Document()
    paragraph = document.add_paragraph(style="Heading 2")
    paragraph.add_run("Signature: ")
    paragraph.add_run("01.10.").bold = True
    paragraph.add_run("2026").italic = True
    paragraph.add_run("; rent: 09.10.2026.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run("Package: weekend").bold = True
    document.sections[0].header.paragraphs[0].text = "Contract series: REF-1"
    document.sections[0].footer.paragraphs[0].text = "Member: no"
    image = io.BytesIO()
    Image.new("RGB", (2, 2), "red").save(image, "PNG")
    image.seek(0)
    document.add_picture(image, width=Inches(0.2))
    document.add_paragraph("Date repeated: Date repeated")
    document.add_paragraph("Line\tbreak")
    document.add_paragraph("Overlapping: aaaa")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def targeted(client, original, replacements, **overrides):
    targets = client.get(f"/api/files/{original['id']}/edit-targets").json()
    changes = []
    for containing, before, after in replacements:
        location = next(item["location"] for item in targets["targets"] if containing in item["text"])
        changes.append({"location": location, "before": before, "after": after})
    return client.post(f"/api/files/{original['id']}/targeted-edit-proposals",
                       json={"before_sha256": original["sha256"], "changes": changes, **overrides})


def test_targeted_docx_edits_split_runs_tables_headers_and_footers_without_rebuilding(client):
    from docx import Document
    from lxml import etree

    source = contract_docx()
    original = upload(client, "test-contract.docx", source)
    targets = client.get(f"/api/files/{original['id']}/edit-targets").json()
    assert any(target["in_table"] and "Package" in target["text"] for target in targets["targets"])
    assert {target["section"] for target in targets["targets"]} == {"body", "header", "footer"}
    response = targeted(client, original, [
        ("Signature:", "01.10.2026", "02.11.2099 (TEST)"),
        ("Signature:", "09.10.2026", "12.11.2099 (TEST)"),
        ("Package:", "weekend", "three days (TEST)"),
        ("Contract series:", "REF-1", "REF-TEST"),
        ("Member:", "no", "yes (TEST)"),
    ])
    assert response.status_code == 200, response.text
    proposal = response.json()
    assert proposal["edit_mode"] == "targeted_text"
    assert len(proposal["changes"]) == 5
    assert all(change["matches"] == 1 for change in proposal["changes"])
    assert "Before:" in proposal["preview_content"] and "After:" in proposal["preview_content"]
    assert "first affected text run" in proposal["warning"]

    preview = client.get(f"/api/files/edit-proposals/{proposal['id']}/preview")
    assert preview.status_code == 200
    assert hashlib.sha256(preview.content).hexdigest() == proposal["after_sha256"]
    assert client.get(f"/api/files/{original['id']}/content").content == source
    assert confirm(client, proposal["id"]).status_code == 200
    updated_bytes = client.get(f"/api/files/{original['id']}/content").content
    assert updated_bytes == preview.content
    updated = Document(io.BytesIO(updated_bytes))
    paragraph = updated.paragraphs[0]
    assert paragraph.style.name == "Heading 2"
    assert paragraph.runs[1].text == "02.11.2099 (TEST)" and paragraph.runs[1].bold
    assert paragraph.runs[2].text == "" and paragraph.runs[2].italic
    assert paragraph.runs[3].text == "; rent: 12.11.2099 (TEST)."
    assert updated.tables[0].cell(0, 0).text == "Package: three days (TEST)"
    assert updated.tables[0].cell(0, 0).paragraphs[0].runs[0].bold
    assert updated.sections[0].header.paragraphs[0].text == "Contract series: REF-TEST"
    assert updated.sections[0].footer.paragraphs[0].text == "Member: yes (TEST)"
    assert len(updated.inline_shapes) == 1
    with zipfile.ZipFile(io.BytesIO(source)) as before, zipfile.ZipFile(io.BytesIO(updated_bytes)) as after:
        assert before.namelist() == after.namelist()
        for part in before.namelist():
            if part not in {"word/document.xml", "word/header1.xml", "word/footer1.xml"}:
                assert before.read(part) == after.read(part)
            else:
                trees = [etree.fromstring(archive.read(part)) for archive in (before, after)]
                for tree in trees:
                    for node in tree.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
                        node.text = ""
                        node.attrib.pop("{http://www.w3.org/XML/1998/namespace}space", None)
                assert etree.tostring(trees[0], method="c14n") == etree.tostring(trees[1], method="c14n")
    assert confirm(client, proposal["id"]).status_code == 409
    assert client.get(f"/api/files/edit-proposals/{proposal['id']}/preview").status_code == 409



def test_legacy_targeted_proposal_scope_still_deserializes_after_upgrade(client, app):
    from pixel_station.file_edits import _digest, _public_proposal

    original = upload(client, "legacy.docx", contract_docx())
    response = targeted(client, original, [("Signature:", "01.10.2026", "02.10.2099")])
    assert response.status_code == 200
    proposal = response.json()
    with app.state.database.session() as session:
        row = session.get(FileEditProposal, proposal["id"])
        row.scope = LEGACY_TARGETED_SCOPE
        row.digest = _digest(row)
        session.commit()
        reloaded = _public_proposal(row)
    assert reloaded["edit_mode"] == "targeted_text"
    assert reloaded["changes"] == proposal["changes"]
    assert reloaded["checkbox_changes"] == []
    assert client.get(f"/api/files/edit-proposals/{proposal['id']}/preview").status_code == 200
    assert confirm(client, proposal["id"]).status_code == 200

@pytest.mark.parametrize("replacement,expected", [
    (("Date repeated", "Date", "TEST"), "found 2"),
    (("Signature:", "unknown date", "TEST"), "found 0"),
    (("Line", "Line\tbreak", "TEST"), "tab or line break"),
    (("Signature:", "01.10.2026", "two\nlines"), "inline text"),
    (("Overlapping:", "aaa", "TEST"), "found 2"),
])
def test_targeted_replacements_reject_ambiguous_missing_or_structural_spans(client, app, replacement, expected):
    source = contract_docx()
    original = upload(client, "test-contract.docx", source)
    response = targeted(client, original, [replacement])
    assert response.status_code == 422
    assert expected in response.text
    assert client.get(f"/api/files/{original['id']}/content").content == source
    with app.state.database.session() as session:
        assert list(session.scalars(select(FileEditProposal))) == []


def test_targeted_replacements_reject_overlap_and_stale_snapshot(client):
    original = upload(client, "test-contract.docx", contract_docx())
    response = targeted(client, original, [("Signature:", "Signature:", "Test"),
                                          ("Signature:", "Signature: 01.10.2026", "Test")])
    assert response.status_code == 422 and "overlap" in response.text
    response = targeted(client, original, [("Signature:", "01.10.2026", "TEST")], before_sha256="0" * 64)
    assert response.status_code == 409
    assert "Reload" in response.text


@pytest.mark.parametrize("alteration", ["original", "stage", "proposal", "expire"])
def test_targeted_preview_and_confirmation_revalidate_reviewed_state(client, app, alteration):
    original = upload(client, "test-contract.docx", contract_docx())
    proposal = targeted(client, original, [("Signature:", "01.10.2026", "TEST")]).json()
    if alteration == "original":
        Path(original["path"]).write_bytes(b"changed externally")
    elif alteration == "stage":
        app.state.file_edits._stage(proposal["id"], ".docx").write_bytes(b"changed stage")
    else:
        with app.state.database.session() as session:
            row = session.get(FileEditProposal, proposal["id"])
            if alteration == "proposal":
                row.content = "altered changes"
            else:
                row.expires_at = 0
            session.commit()
    assert client.get(f"/api/files/edit-proposals/{proposal['id']}/preview").status_code == 409
    assert confirm(client, proposal["id"]).status_code == 409


def test_targeted_docx_rejects_paragraph_with_word_fields(client):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    paragraph = document.add_paragraph("Signature: 01.10.2026")
    field = OxmlElement("w:fldChar")
    field.set(qn("w:fldCharType"), "begin")
    paragraph.add_run()._r.append(field)
    output = io.BytesIO()
    document.save(output)
    original = upload(client, "field-contract.docx", output.getvalue())
    response = targeted(client, original, [("Signature:", "01.10.2026", "TEST")])
    assert response.status_code == 422 and "Fields or tracked changes" in response.text


@pytest.mark.asyncio
async def test_chat_docx_uses_small_targeted_changes_and_keeps_original(app):
    from pixel_station.files import ingest

    class TargetedLLM:
        async def structured(self, model, messages, schema, **kwargs):
            assert schema is PlannedTargetedEdit
            assert "Do not regenerate the document" in messages[0]["content"]
            assert "Do not invent missing dates" in messages[0]["content"]
            return schema(file_id=original.id, plan="Use the explicit TEST signature date.",
                          changes=[{"location": "word/document.xml:p:0", "before": "01.10.2026", "after": "02.11.2099 (TEST)"}])

    source = contract_docx()
    with app.state.database.session() as session:
        original = ingest(session, app.state.data_dir, "test-contract.docx", source)
    app.state.llm = TargetedLLM()
    settings = app.state.settings()
    settings.roles["planner"] = "test-local"
    app.state.set_settings(settings)
    result = await propose_chat_edit(app, [original.id], "Set the signature date to 02.11.2099 (TEST).")
    assert result["file_edit"]["edit_mode"] == "targeted_text"
    assert result["file_edit"]["changes"][0]["matches"] == 1
    assert Path(original.path).read_bytes() == source


def test_named_library_copy_has_own_record_chunks_and_revisions_without_changing_source(client, app):
    source = contract_docx()
    original = upload(client, "master.docx", source)
    original_record = client.get(f"/api/files/{original['id']}").json()
    metadata = Path(original["path"]).parent / "records" / original["id"] / "metadata.json"
    original_metadata = metadata.read_bytes()
    response = client.post(f"/api/files/{original['id']}/copy", json={"filename": "TEST client copy.docx"})
    assert response.status_code == 200, response.text
    copy = response.json()
    assert copy["id"] != original["id"]
    assert copy["filename"] == "TEST client copy.docx" and copy["source"] == "copied"
    assert copy["sha256"] == original["sha256"] and copy["path"] == original["path"]
    copy_record = client.get(f"/api/files/{copy['id']}").json()
    assert {chunk["id"] for chunk in copy_record["chunks"]}.isdisjoint(chunk["id"] for chunk in original_record["chunks"])
    assert [chunk["text"] for chunk in copy_record["chunks"]] == [chunk["text"] for chunk in original_record["chunks"]]
    proposal = targeted(client, copy, [("Signature:", "01.10.2026", "TEST date")]).json()
    assert confirm(client, proposal["id"]).status_code == 200
    assert client.get(f"/api/files/{original['id']}/content").content == source
    assert client.get(f"/api/files/{original['id']}").json() == original_record
    assert metadata.read_bytes() == original_metadata
    assert len(client.get(f"/api/files/{copy['id']}/revisions").json()) == 1
    assert client.get(f"/api/files/{original['id']}/revisions").json() == []
    with app.state.database.session() as session:
        assert session.get(Attachment, original["id"]).sha256 == original["sha256"]


def test_library_copy_rejects_changed_format_or_external_bytes(client):
    original = upload(client)
    assert client.post(f"/api/files/{original['id']}/copy", json={"filename": "invalid.pdf"}).status_code == 422
    Path(original["path"]).write_bytes(b"external mutation")
    response = client.post(f"/api/files/{original['id']}/copy", json={"filename": "copy.txt"})
    assert response.status_code == 409
