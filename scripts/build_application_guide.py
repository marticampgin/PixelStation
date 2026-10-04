"""Build the public application guide with the bundled document runtime.

Run this file with the bundled document Python runtime;
no application data, credentials or private examples are read by the builder.
"""

import argparse
import html
import re
from pathlib import Path
from urllib.parse import urljoin

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / "docs" / "design" / "application-guide"
INK, ACCENT, PALE, BORDER = "172844", "6351A5", "F3F5FB", "D9D9D9"
DIAGRAMS = {
    "components": {
        "nodes": [
            (1, 0, "Browser interface", "React and TypeScript\nPort 5173"),
            (1, 1, "Application code", "Routes and permissions\nPython at port 8000"),
            (0, 2, "Local model runtime", "Ollama at 11434\nText vision embeddings"),
            (1, 2, "Local saved data", "SQLite files revisions\nMemory runs and jobs"),
            (2, 2, "Local image runtime", "ComfyUI at 8188\nSaved image graph"),
            (0, 3, "Web research", "SearXNG at 8888\nInternet search and pages"),
            (2, 3, "Google services", "Authorized Gmail\nSeparate Calendar account"),
        ],
        "edges": [(0, 1), (1, 2), (1, 3), (1, 4), (2, 5), (4, 6)],
        "edge_override": {(2, 5): (1, 5), (4, 6): (1, 6)},
        "alt": "The browser calls Python. Python controls local model, data, image and authorized internet services.",
    },
    "request": {
        "nodes": [
            (0, 0, "Save request", "Message and run record"),
            (1, 0, "Choose workflow", "Rules first\nClassifier if ambiguous"),
            (2, 0, "Gather evidence", "Relevant memory files\nSelected sources"),
            (2, 1, "Execute tools", "Validated arguments\nBounded permissions"),
            (1, 1, "Answer and validate", "Public text only\nOne bounded repair"),
            (0, 1, "Stream and persist", "Text progress artifacts\nOutcome and measurements"),
        ],
        "edges": [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5)],
        "alt": "A request is saved, routed, given evidence and tools, answered and validated, then streamed and saved.",
    },
    "memory": {
        "nodes": [
            (0, 0, "Stored knowledge", "Memories chunks\nConversation summaries"),
            (1, 0, "Retrieve candidates", "Word match and vectors\nScope pins importance age"),
            (2, 0, "Select bounded context", "Four memories by default\nDocument excerpts"),
            (2, 1, "Build answer input", "Newest request retained\nOlder history clipped"),
            (1, 1, "Answer locally", "Context is supplied text\nWeights are not retrained"),
            (0, 1, "Idle background work", "Summarize user turns\nExtract and index candidates"),
        ],
        "edges": [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5), (5, 0)],
        "alt": "Retrieval selects a small context from saved records. Idle work later creates summaries and indexes eligible user memories.",
    },
    "files": {
        "nodes": [
            (0, 0, "Upload or copy", "Immutable byte fingerprint\nIndependent named record"),
            (1, 0, "Parse and index", "Local format readers\nOffline OCR when needed"),
            (2, 0, "Retrieve or propose", "Cited excerpts\nExact reviewed replacement"),
            (2, 1, "Inspect preview", "Original hash and changes\nExpiry and permission"),
            (1, 1, "Confirm once", "Verify reviewed bytes\nPreserve prior revision"),
            (0, 1, "Updated library copy", "Reparse and reindex\nHost original untouched"),
        ],
        "edges": [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5)],
        "alt": "File ingestion stores and indexes a record. Editing requires a reviewed preview and one confirmation before revision preservation and reindexing.",
    },
    "approval": {
        "nodes": [
            (0, 0, "Requested action", "Send edit or Calendar\nValidated fields"),
            (1, 0, "Bind exact proposal", "Payload byte hashes\nCurrent account or record"),
            (2, 0, "Human review", "Recipient terms times\nFiles and visible changes"),
            (2, 1, "Confirm once", "Within ten minutes\nBinding still unchanged"),
            (1, 1, "Execute and verify", "Returned remote identifier\nReopen where supported"),
            (0, 1, "Persist outcome", "Accepted or failed\nNo blind resubmission"),
        ],
        "edges": [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5)],
        "alt": "Consequential actions bind an exact proposal, require human review and single-use confirmation, then execute and retain verification outcome.",
    },
    "resources": {
        "nodes": [
            (0, 0, "Shared queue lock", "Chat Poker embeddings\nImages and idle model work"),
            (1, 0, "Observe Ollama", "Fresh loaded model list\nUnload and verify absence"),
            (2, 0, "Run ComfyUI job", "Bound workflow endpoint\nProgress history output"),
            (2, 1, "Confirm job inactive", "Owned prompt absent\nCancellation also checked"),
            (1, 1, "Request release", "Free models and memory\nRecord measured counters"),
            (0, 1, "Release shared lock", "Retain image and warnings\nNext owned work may run"),
        ],
        "edges": [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5)],
        "alt": "The image lifecycle holds the shared queue, verifies local text model unloading, executes its job, confirms remote inactivity and checks memory release.",
    },
    "watchtower": {
        "nodes": [
            (0, 0, "Observe real runs", "Timing tokens outcomes\nRepairs user feedback"),
            (1, 0, "Group exact evidence", "Route model error stage\nContent free linked runs"),
            (2, 0, "Define a regression", "Fictional input\nExplicit expected behavior"),
            (2, 1, "Run isolated checks", "Temporary data\nNo real mail or events"),
            (1, 1, "Opt in to native probes", "Actual installed model\nSeparate quality criteria"),
            (0, 1, "Review and change", "One change at a time\nKeep failed observations"),
        ],
        "edges": [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5)],
        "alt": "Observed failures lead to explicit regressions and optional native model probes. Results inform reviewed changes without automatic source editing.",
    },
}


def font(size, bold=False):
    candidates = [
        Path("C:/Windows/Fonts") / ("segoeuib.ttf" if bold else "segoeui.ttf"),
        Path("/usr/share/fonts/truetype/dejavu") / ("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"),
    ]
    return ImageFont.truetype(str(next(path for path in candidates if path.exists())), size)


def render_diagrams():
    ART.mkdir(parents=True, exist_ok=True)
    for name, spec in DIAGRAMS.items():
        width, box_w, box_h, gx, gy, margin = 1800, 500, 230, 75, 110, 75
        rows = max(node[1] for node in spec["nodes"]) + 1
        height = rows * box_h + (rows - 1) * gy + margin * 2
        picture = Image.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(picture)
        boxes = [(margin + x * (box_w + gx), margin + y * (box_h + gy)) for x, y, *_ in spec["nodes"]]
        svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
               '<rect width="100%" height="100%" fill="white"/>',
               '<defs><marker id="a" markerWidth="12" markerHeight="12" refX="9" refY="5" orient="auto"><path d="M0 0 L10 5 L0 10" fill="#6351A5"/></marker></defs>']
        for edge in spec["edges"]:
            source, target = spec.get("edge_override", {}).get(edge, edge)
            sx, sy = boxes[source]
            tx, ty = boxes[target]
            if sy == ty:
                start = (sx + (box_w if tx > sx else 0), sy + box_h // 2)
                end = (tx + (0 if tx > sx else box_w), ty + box_h // 2)
            else:
                start = (sx + box_w // 2, sy + (box_h if ty > sy else 0))
                end = (tx + box_w // 2, ty + (0 if ty > sy else box_h))
            points = [start, end]
            if name == "components" and target in {5, 6}:
                outward = target == 6
                start = (sx + (box_w if outward else 0), sy + box_h // 2)
                end = (tx + (box_w if outward else 0), ty + box_h // 2)
                outer_x = width - 30 if outward else 30
                points = [start, (outer_x, start[1]), (outer_x, end[1]), end]
            draw.line(points, fill="#" + ACCENT, width=5)
            import math
            previous = points[-2]
            angle = math.atan2(end[1] - previous[1], end[0] - previous[0])
            tip = [end, (end[0] - 22 * math.cos(angle - .4), end[1] - 22 * math.sin(angle - .4)),
                   (end[0] - 22 * math.cos(angle + .4), end[1] - 22 * math.sin(angle + .4))]
            draw.polygon(tip, fill="#" + ACCENT)
            coordinates = " ".join(f"{x},{y}" for x, y in points)
            svg.append(f'<polyline points="{coordinates}" fill="none" stroke="#{ACCENT}" stroke-width="5" marker-end="url(#a)"/>')
        for index, (_, _, title, body) in enumerate(spec["nodes"]):
            x, y = boxes[index]
            draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=14, fill="#" + PALE, outline="#" + BORDER, width=3)
            draw.text((x + 25, y + 27), title, font=font(38, True), fill="#" + INK)
            for row, line in enumerate(body.splitlines()):
                draw.text((x + 25, y + 94 + row * 48), line, font=font(35), fill="#" + INK)
            svg.append(f'<rect x="{x}" y="{y}" width="{box_w}" height="{box_h}" rx="14" fill="#{PALE}" stroke="#{BORDER}" stroke-width="3"/>')
            svg.append(f'<text x="{x+25}" y="{y+65}" font-family="Segoe UI,Arial,sans-serif" font-size="38" font-weight="bold" fill="#{INK}">{html.escape(title)}</text>')
            for row, line in enumerate(body.splitlines()):
                svg.append(f'<text x="{x+25}" y="{y+129+row*48}" font-family="Segoe UI,Arial,sans-serif" font-size="35" fill="#{INK}">{html.escape(line)}</text>')
        picture.save(ART / f"{name}.png", dpi=(300, 300))
        (ART / f"{name}.svg").write_text("\n".join(svg + ["</svg>"]), encoding="utf-8")


def inline(paragraph, text):
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^]]+\]\([^)]+\))", text):
        if match := re.fullmatch(r"\[([^]]+)\]\(([^)]+)\)", part):
            label, target = match.groups()
            link = OxmlElement("w:hyperlink")
            target = urljoin("https://github.com/marticampgin/PixelStation/blob/main/docs/", target)
            relation = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            link.set(qn("r:id"), relation)
            run, props = OxmlElement("w:r"), OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), ACCENT)
            props.append(color)
            run.append(props)
            value = OxmlElement("w:t")
            value.text = label
            run.append(value)
            link.append(run)
            paragraph._p.append(link)
            continue
        run = paragraph.add_run(part[1:-1] if part.startswith("`") else part[2:-2] if part.startswith("**") else part)
        if part.startswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(10.5)
        elif part.startswith("**"):
            run.bold = True


def bookmark(paragraph, name, number):
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), name)
    end.set(qn("w:id"), str(number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def toc_link(paragraph, title, target):
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), target)
    run, props = OxmlElement("w:r"), OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT)
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = title
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def table(document, rows):
    cells = [re.split(r"(?<!\\)\|", row.strip().strip("|")) for row in rows]
    if len(cells) > 1 and all(re.fullmatch(r"\s*:?-+:?\s*", cell) for cell in cells[1]):
        cells.pop(1)
    count = len(cells[0])
    tab = document.add_table(rows=0, cols=count)
    tab.autofit = False
    widths = {2: [2.1, 4.8], 3: [1.55, 2.1, 3.25], 4: [1.4, 1.2, 2.9, 1.4]}.get(count, [6.9 / count] * count)
    for column, width in zip(tab.columns, widths, strict=True):
        column.width = Inches(width)
    for index, values in enumerate(cells):
        row = tab.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, value, width in zip(row.cells, values, widths, strict=True):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement("w:" + edge)
                for key, val in {"val": "single", "sz": "4", "color": BORDER}.items():
                    border.set(qn("w:" + key), val)
                borders.append(border)
            tc.append(borders)
            fill = OxmlElement("w:shd")
            fill.set(qn("w:fill"), INK if index == 0 else PALE if index % 2 == 0 else "FFFFFF")
            tc.append(fill)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "bottom", "left", "right"):
                part = OxmlElement("w:" + edge)
                part.set(qn("w:w"), "95")
                part.set(qn("w:type"), "dxa")
                margins.append(part)
            tc.append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.space_before = Pt(3)
            inline(paragraph, value.strip().replace("\\|", "|"))
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                if index == 0:
                    run.font.color.rgb = RGBColor.from_string("FFFFFF")
                    run.bold = True
        if index == 0:
            header = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(header)
    document.add_paragraph().paragraph_format.space_after = Pt(3)


def build(source, output):
    render_diagrams()
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.8)
    section.top_margin = section.bottom_margin = Inches(.75)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11.5)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(7)
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor.from_string("000000")
        style.font.italic = False
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(16 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(8)
    document.styles["Title"].font.size = Pt(30)
    document.styles["Heading 1"].font.size = Pt(19)
    document.styles["Heading 2"].font.size = Pt(15)
    document.styles["Heading 3"].font.size = Pt(12)
    document.core_properties.title = "Know Your Pixel Station"
    document.core_properties.subject = "Application architecture settings operational reference and concrete workflows"
    document.core_properties.author = "Pixel Station"
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Pixel Station  |  ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    content = source.read_text(encoding="utf-8")
    reference = source.parent / "APPLICATION_REFERENCE.md"
    if reference.exists():
        content = content.replace("<!-- GENERATED_REFERENCE -->", reference.read_text(encoding="utf-8"))
    lines = content.splitlines()
    headings = [line[3:] for line in lines if line.startswith("## ")]
    anchors = {title: f"section_{i}" for i, title in enumerate(headings)}
    i, number, toc_done = 0, 0, False
    while i < len(lines):
        line = lines[i]
        if not line.strip() or line.startswith("<!--"):
            i += 1
            continue
        if line.startswith("# "):
            if i == 0:
                document.add_paragraph(line[2:], "Title")
                document.add_paragraph("Architecture workflows settings and operational reference", "Subtitle")
                document.add_paragraph("Installation snapshot 4 October 2026  |  Version 0.1.0")
        elif line.startswith("## "):
            title = line[3:]
            if not toc_done and title.startswith("1 "):
                document.add_heading("Contents", level=1)
                for text in headings:
                    if text != "What this guide explains":
                        toc_link(document.add_paragraph(), text, anchors[text])
                document.add_page_break()
                toc_done = True
            paragraph = document.add_heading(title, level=1)
            bookmark(paragraph, anchors[title], number)
            number += 1
        elif line.startswith("### "):
            document.add_heading(line[4:], level=2)
        elif line.startswith("#### "):
            document.add_heading(line[5:], level=3)
        elif match := re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line):
            label, target = match.groups()
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            run = paragraph.add_run()
            run.add_picture(str(source.parent / target), width=Inches(6.9))
            drawing = run._r.xpath(".//wp:docPr")[0]
            drawing.set("descr", DIAGRAMS.get(Path(target).stem, {}).get("alt", label))
            caption = document.add_paragraph(label)
            caption.paragraph_format.space_after = Pt(10)
            caption.runs[0].italic = True
            caption.runs[0].font.size = Pt(10)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            table(document, rows)
            continue
        elif line.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                paragraph = document.add_paragraph()
                run = paragraph.add_run(lines[i])
                run.font.name, run.font.size = "Consolas", Pt(9.5)
                paragraph.paragraph_format.space_after = Pt(2)
                i += 1
        else:
            style = "List Bullet" if line.startswith("- ") else "List Number" if re.match(r"^\d+\. ", line) else "Normal"
            value = line[2:] if style == "List Bullet" else re.sub(r"^\d+\. ", "", line) if style == "List Number" else line
            inline(document.add_paragraph(style=style), value)
        i += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Saved {output}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=ROOT / "docs" / "KNOW_YOUR_APPLICATION.md")
    parser.add_argument("--output", type=Path, default=ROOT / "data" / "exports" / "Know Your Pixel Station.docx")
    arguments = parser.parse_args()
    build(arguments.source, arguments.output)
